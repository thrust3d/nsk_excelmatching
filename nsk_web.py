#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
app.py
======
Web sučelje (Streamlit) za provjeru_duplikata.py.

Korisnik može:
  - povući/odabrati više .xlsx datoteka,
  - po želji podesiti pragove (skriveno pod "Napredne postavke"),
  - kliknuti "Pokreni provjeru",
  - preuzeti izvjestaj_duplikati.xlsx.

Sva logika prepoznavanja duplikata je IDENTIČNA originalnom
provjera_duplikata.py skriptom - ovdje je samo dodano web sučelje
oko nje (CLI argparse je zamijenjen Streamlit kontrolama).
"""

import os
import re
import io
import sys
import shutil
import tempfile
import unicodedata
from itertools import combinations
from collections import defaultdict

import pandas as pd
import streamlit as st
from rapidfuzz import fuzz
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

# ==========================================================================
# ISTI PRAGOVI I LOGIKA KAO U ORIGINALNOJ SKRIPTI (provjera_duplikata.py)
# ==========================================================================
PRAG_NASLOV_MIN   = 0.55
PRAG_PRIJAVE_DEF  = 0.55
PRAG_GRUPA_DEF    = 0.90
PRAG_VISOK        = 0.80
PRAG_SERIJA_ISTI  = 0.90
MIN_TOKEN_LEN     = 3

KAND_NASLOV = ["naslov", "title"]
KAND_AUTOR  = ["autor", "podaci o odgovornosti", "author"]
KAND_GODINA = ["godina izdavanja", "godina", "razdoblje izdavanja", "razdoblje  izdavanja", "year"]

RIJECI_ODGOVORNOSTI = {
    "slozio", "uredio", "priredio", "napisao", "napisala", "sastavio",
    "preveo", "prevela", "izdao", "izdala", "uredila", "priredila",
    "odgovorni", "urednik", "ur", "nakl", "naklada", "tekst", "autor",
}


def skini_dijakritike(s: str) -> str:
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    return s.replace("đ", "d").replace("Đ", "d")


def norm_naslov(s) -> str:
    if s is None:
        return ""
    s = str(s)
    if s.strip().lower() in ("nan", "/", "\xa0", ""):
        return ""
    s = s.split(" / ")[0]
    s = skini_dijakritike(s).lower()
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s


def norm_autor(s) -> str:
    if s is None:
        return ""
    s = str(s)
    if s.strip().lower() in ("nan", "/", "\xa0", ""):
        return ""
    s = skini_dijakritike(s).lower()
    s = re.sub(r"\[.*?\]", " ", s)
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    toks = [t for t in s.split() if t not in RIJECI_ODGOVORNOSTI]
    return " ".join(toks).strip()


def parsiraj_godinu(s):
    if s is None:
        return None
    g = re.findall(r"(1[5-9]\d{2}|20\d{2})", str(s))
    if not g:
        return None
    return min(int(x) for x in g)


def tokeni(norm_str):
    return {t for t in norm_str.split() if len(t) >= MIN_TOKEN_LEN}


def trigrami(norm_str):
    s = norm_str.replace(" ", "")
    return {s[i:i + 3] for i in range(len(s) - 2)} if len(s) >= 3 else {s}


def nadji_stupac(stupci, kandidati):
    norm_map = {skini_dijakritike(str(c)).lower().strip(): c for c in stupci}
    for k in kandidati:
        kk = skini_dijakritike(k).lower().strip()
        if kk in norm_map:
            return norm_map[kk]
    for k in kandidati:
        kk = skini_dijakritike(k).lower().strip()
        for nm, orig in norm_map.items():
            if kk and (kk in nm or nm in kk):
                return orig
    return None


def list_je_serijski(naziv_lista, stupci):
    nl = skini_dijakritike(str(naziv_lista)).lower()
    if "serij" in nl:
        return True
    for c in stupci:
        if "razdoblje" in skini_dijakritike(str(c)).lower():
            return True
    return False


def ucitaj_unose_iz_putanja(putanje):
    """Ista logika kao ucitaj_unose() iz originalne skripte, ali prima
    eksplicitnu listu putanja do .xlsx/.xlsm datoteka (npr. privremeno
    spremljene uploadane datoteke), umjesto da pretražuje folder."""
    unosi = []
    upozorenja = []
    for path in putanje:
        fname = os.path.basename(path)
        try:
            xl = pd.ExcelFile(path)
        except Exception as e:
            upozorenja.append(f"Preskačem '{fname}': {e}")
            continue
        for sheet in xl.sheet_names:
            df = pd.read_excel(path, sheet_name=sheet, header=0, dtype=str)
            c_nas = nadji_stupac(df.columns, KAND_NASLOV)
            c_aut = nadji_stupac(df.columns, KAND_AUTOR)
            c_god = nadji_stupac(df.columns, KAND_GODINA)
            serijski = list_je_serijski(sheet, df.columns)
            if c_nas is None:
                upozorenja.append(f"[{fname} / {sheet}] nema stupca naslova -> preskačem.")
                continue
            for idx, row in df.iterrows():
                raw_nas = row.get(c_nas)
                nn = norm_naslov(raw_nas)
                if not nn:
                    continue
                raw_aut = row.get(c_aut) if c_aut else None
                raw_god = row.get(c_god) if c_god else None
                unosi.append({
                    "datoteka": fname, "list": sheet, "redak_excel": idx + 2,
                    "serijski": serijski,
                    "naslov_raw": "" if raw_nas is None else str(raw_nas),
                    "autor_raw":  "" if raw_aut is None else str(raw_aut),
                    "godina_raw": "" if raw_god is None else str(raw_god),
                    "n_naslov": nn, "n_autor": norm_autor(raw_aut),
                    "godina": parsiraj_godinu(raw_god),
                    "tokeni": tokeni(nn), "trig": trigrami(nn),
                })
    return unosi, upozorenja


def kandidat_parovi(unosi):
    inv = defaultdict(list)
    for i, u in enumerate(unosi):
        kljucevi = u["tokeni"] | {"§" + t for t in u["trig"]}
        if not kljucevi:
            kljucevi = {"§§short"}
        for k in kljucevi:
            inv[k].append(i)
    parovi = set()
    for idxs in inv.values():
        if len(idxs) > 1:
            for a, b in combinations(sorted(idxs), 2):
                parovi.add((a, b))
    return parovi


def naslov_slicnost(a, b):
    ts = fuzz.token_set_ratio(a["n_naslov"], b["n_naslov"]) / 100.0
    sa, sb = a["tokeni"], b["tokeni"]
    if sa and sb:
        manji, veci = (sa, sb) if len(sa) <= len(sb) else (sb, sa)
        contain = len(manji & veci) / len(manji)
    else:
        contain = 0.0
    return max(ts, contain)


def autor_slicnost(a, b):
    if not a["n_autor"] or not b["n_autor"]:
        return None
    return fuzz.token_set_ratio(a["n_autor"], b["n_autor"]) / 100.0


def grupna_slicnost(a, b):
    return fuzz.token_sort_ratio(a["n_naslov"], b["n_naslov"]) / 100.0


def ocijeni(a, b, prag_prijave):
    t = naslov_slicnost(a, b)
    if t < PRAG_NASLOV_MIN:
        return None
    aut = autor_slicnost(a, b)
    g1, g2 = a["godina"], b["godina"]
    god_dostupna = g1 is not None and g2 is not None
    god_delta = abs(g1 - g2) if god_dostupna else None

    if aut is not None:
        if t >= PRAG_VISOK and aut >= PRAG_VISOK:
            conf = min(0.99, 0.5 * t + 0.5 * aut + 0.05)
        else:
            conf = 0.60 * t + 0.40 * aut
            if god_dostupna:
                if god_delta == 0:
                    conf = min(0.99, conf + 0.08)
                elif god_delta > 2:
                    conf = max(0.5 * t, conf - 0.08)
        nap_aut = f"{aut*100:.0f}%"
    else:
        conf = t
        if god_dostupna:
            if god_delta == 0:
                conf = min(0.99, conf + 0.10)
            elif god_delta > 2:
                conf = max(0.5 * t, conf - 0.10)
        nap_aut = "n/d"

    if conf < prag_prijave:
        return None
    god_opis = "n/d" if not god_dostupna else ("ista" if god_delta == 0 else f"Δ {god_delta} god.")
    return {"conf": conf, "naslov_pct": t, "autor_opis": nap_aut, "godina_opis": god_opis}


def serijski_preskoci(a, b):
    if a["serijski"] and b["serijski"]:
        if naslov_slicnost(a, b) >= PRAG_SERIJA_ISTI:
            return True
    return False


class UF:
    def __init__(self, n):
        self.p = list(range(n))

    def find(self, x):
        while self.p[x] != x:
            self.p[x] = self.p[self.p[x]]
            x = self.p[x]
        return x

    def union(self, a, b):
        ra, rb = self.find(a), self.find(b)
        if ra != rb:
            self.p[ra] = rb


THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _stiliziraj_zaglavlje(ws):
    fill = PatternFill("solid", fgColor="1F4E78")
    font = Font(bold=True, color="FFFFFF", name="Arial", size=10)
    for c in ws[1]:
        c.fill = fill
        c.font = font
        c.border = BORDER
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def napisi_izvjestaj_u_memoriju(grupe, rubni, unosi, broj_datoteka, prag_prijave, prag_grupa):
    wb = Workbook()

    ws = wb.active
    ws.title = "Grupe"
    ws.append(["Grupa #", "Veličina", "Sličnost u grupi %", "Naslov", "Autor",
               "Godina", "Datoteka", "List", "Redak"])
    _stiliziraj_zaglavlje(ws)

    paleta = ["FCE4D6", "DDEBF7", "E2EFDA", "FFF2CC", "EAD1DC", "D9D9D9"]
    for gi, clanovi in enumerate(grupe, start=1):
        rep = max(clanovi, key=lambda m: len(unosi[m]["n_naslov"]))
        boja = paleta[(gi - 1) % len(paleta)]
        red_clanovi = [rep] + [m for m in clanovi if m != rep]
        for pos, m in enumerate(red_clanovi):
            u = unosi[m]
            if m == rep:
                pct = "predstavnik"
            else:
                pct = round(grupna_slicnost(u, unosi[rep]) * 100)
            ws.append([
                gi if pos == 0 else "",
                len(clanovi) if pos == 0 else "",
                pct, u["naslov_raw"], u["autor_raw"], u["godina_raw"],
                u["datoteka"], u["list"], u["redak_excel"],
            ])
            for c in ws[ws.max_row]:
                c.fill = PatternFill("solid", fgColor=boja)
                c.font = Font(name="Arial", size=10, bold=(pos == 0))
                c.border = BORDER
                c.alignment = Alignment(vertical="top", wrap_text=True)
    sirine = [9, 9, 16, 40, 24, 9, 26, 20, 8]
    for i, w in enumerate(sirine, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wp = wb.create_sheet("Rubne veze")
    wp.append(["Sličnost %", "Naslov %", "Autor", "Godina", "Predst. parova",
               "Naslov A", "Čl. A", "Autor A", "Godina A", "Datoteka A", "List A", "Redak A",
               "Naslov B", "Čl. B", "Autor B", "Godina B", "Datoteka B", "List B", "Redak B"])
    _stiliziraj_zaglavlje(wp)
    for i, j, sc, cnt, szA, szB in rubni:
        a, b = unosi[i], unosi[j]
        wp.append([
            round(sc["conf"] * 100), round(sc["naslov_pct"] * 100),
            sc["autor_opis"], sc["godina_opis"], cnt,
            a["naslov_raw"], szA, a["autor_raw"], a["godina_raw"], a["datoteka"], a["list"], a["redak_excel"],
            b["naslov_raw"], szB, b["autor_raw"], b["godina_raw"], b["datoteka"], b["list"], b["redak_excel"],
        ])
    for row in wp.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=10)
            c.border = BORDER
            c.alignment = Alignment(vertical="top", wrap_text=True)
        row[0].alignment = Alignment(horizontal="center", vertical="top")
    for i, w in enumerate([10, 9, 8, 9, 9, 32, 6, 20, 8, 20, 16, 7, 32, 6, 20, 8, 20, 16, 7], start=1):
        wp.column_dimensions[wp.cell(row=1, column=i).column_letter].width = w
    wp.freeze_panes = "A2"
    wp.auto_filter.ref = wp.dimensions

    s = wb.create_sheet("Sažetak")
    s["A1"] = "Sažetak provjere duplikata"
    s["A1"].font = Font(bold=True, size=12, name="Arial")
    clanova_u_grupama = sum(len(g) for g in grupe)
    redovi = [
        ("Pregledano datoteka", broj_datoteka),
        ("Pregledano unosa (knjiga)", len(unosi)),
        ("Broj grupa (≥2 člana)", len(grupe)),
        ("Ukupno članova u grupama", clanova_u_grupama),
        ("Najveća grupa", max((len(g) for g in grupe), default=0)),
        ("Rubnih veza (parovi grupa/unosa)", len(rubni)),
        ("Prag grupiranja", f"{prag_grupa*100:.0f}%"),
        ("Prag prijave", f"{prag_prijave*100:.0f}%"),
    ]
    for i, (k, v) in enumerate(redovi, start=3):
        s[f"A{i}"] = k
        s[f"B{i}"] = v
        s[f"A{i}"].font = Font(name="Arial", size=10)
        s[f"B{i}"].font = Font(name="Arial", size=10, bold=True)
    s.column_dimensions["A"].width = 34
    s.column_dimensions["B"].width = 14

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def pokreni_analizu(unosi, prag_prijave, prag_grupa):
    """Cijela logika iz main() originalne skripte, izdvojena u funkciju
    koju poziva web sučelje."""
    parovi = kandidat_parovi(unosi)
    svi = []
    for i, j in parovi:
        a, b = unosi[i], unosi[j]
        if serijski_preskoci(a, b):
            continue
        sc = ocijeni(a, b, prag_prijave)
        if sc:
            svi.append((i, j, sc))

    uf = UF(len(unosi))
    jaki_cvorovi = set()
    grupna = {}
    for i, j, sc in svi:
        g = grupna_slicnost(unosi[i], unosi[j])
        grupna[(i, j)] = g
        if g >= prag_grupa:
            uf.union(i, j)
            jaki_cvorovi.add(i)
            jaki_cvorovi.add(j)

    korijeni = defaultdict(list)
    for idx in jaki_cvorovi:
        korijeni[uf.find(idx)].append(idx)
    grupe = [sorted(v) for v in korijeni.values() if len(v) >= 2]
    grupe.sort(key=len, reverse=True)

    cid = {}
    cluster_size = {}
    for r, members in korijeni.items():
        if len(members) >= 2:
            cluster_size[("g", r)] = len(members)
            for m in members:
                cid[m] = ("g", r)

    def get_cid(idx):
        return cid.get(idx, ("s", idx))

    def csize(c):
        return cluster_size.get(c, 1)

    agg = {}
    for i, j, sc in svi:
        if grupna[(i, j)] >= prag_grupa:
            continue
        ca, cb = get_cid(i), get_cid(j)
        if ca == cb:
            continue
        key = tuple(sorted([ca, cb], key=lambda x: (x[0], str(x[1]))))
        e = agg.get(key)
        if e is None:
            agg[key] = {"best": (i, j, sc), "count": 1}
        else:
            e["count"] += 1
            if sc["conf"] > e["best"][2]["conf"]:
                e["best"] = (i, j, sc)
    rubni = []
    for e in agg.values():
        i, j, sc = e["best"]
        rubni.append((i, j, sc, e["count"], csize(get_cid(i)), csize(get_cid(j))))
    rubni.sort(key=lambda t: t[2]["conf"], reverse=True)

    return grupe, rubni


# ==========================================================================
# STREAMLIT SUČELJE
# ==========================================================================
st.set_page_config(page_title="Provjera duplikata knjiga", layout="centered")

st.title("Provjera duplikata knjiga")
st.write(
    "Učitajte jednu ili više Excel datoteka s popisom knjiga. "
    "Alat pretraži sve listove i pronađe knjige koje se vjerojatno "
    "ponavljaju, unutar iste datoteke ili između više njih."
)

uploaded_files = st.file_uploader(
    "Excel datoteke (.xlsx)",
    type=["xlsx", "xlsm"],
    accept_multiple_files=True,
)

with st.expander("Napredne postavke"):
    prag_prijave = st.slider(
        "Prag prijave",
        min_value=0.30, max_value=0.95, value=PRAG_PRIJAVE_DEF, step=0.01,
        help="Niža vrijednost = više prijava, uključujući manje sigurne. Zadano: 0.55",
    )
    prag_grupa = st.slider(
        "Prag grupiranja",
        min_value=0.50, max_value=0.99, value=PRAG_GRUPA_DEF, step=0.01,
        help="Koliko naslov mora biti gotovo identičan da se spoji u istu grupu. Zadano: 0.90",
    )

pokreni = st.button("Pokreni provjeru", type="primary", disabled=not uploaded_files)

if pokreni and uploaded_files:
    with st.spinner("Obrađujem datoteke..."):
        tmp_dir = tempfile.mkdtemp()
        try:
            putanje = []
            for uf_ in uploaded_files:
                dest = os.path.join(tmp_dir, uf_.name)
                with open(dest, "wb") as f:
                    f.write(uf_.getbuffer())
                putanje.append(dest)

            unosi, upozorenja = ucitaj_unose_iz_putanja(putanje)
            broj_datoteka = len({u["datoteka"] for u in unosi})

            if not unosi:
                st.error(
                    "Nije pronađen nijedan unos s prepoznatim stupcem naslova. "
                    "Provjerite imaju li tablice stupac 'Naslov' (ili 'Title')."
                )
            else:
                grupe, rubni = pokreni_analizu(unosi, prag_prijave, prag_grupa)
                buf = napisi_izvjestaj_u_memoriju(
                    grupe, rubni, unosi, broj_datoteka, prag_prijave, prag_grupa
                )

                st.write("Provjera završena.")
                col1, col2, col3 = st.columns(3)
                col1.metric("Pregledano knjiga", len(unosi))
                col2.metric("Grupe duplikata", len(grupe))
                col3.metric("Rubne veze", len(rubni))

                st.download_button(
                    label="Preuzmi izvještaj",
                    data=buf,
                    file_name="izvjestaj_duplikati.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    type="primary",
                )

                if upozorenja:
                    with st.expander("Upozorenja pri učitavanju"):
                        for w in upozorenja:
                            st.write("- " + w)
        finally:
            shutil.rmtree(tmp_dir, ignore_errors=True)

st.divider()
st.caption(
    "Datoteke se obrađuju privremeno tijekom analize i brišu se odmah "
    "nakon generiranja izvještaja. Ništa se ne pohranjuje trajno."
)